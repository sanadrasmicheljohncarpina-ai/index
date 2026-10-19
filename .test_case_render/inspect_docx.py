from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

p = Path(__file__).resolve().parents[1] / "docs" / "TEST_CASES_Run_2026-10-06.docx"
d = Document(p)
t = d.tables[0]
body_fills = []
for table in (d.tables[0], d.tables[1]):
    for row in table.rows[1:]:
        for cell in row.cells:
            shd = cell._tc.tcPr.find(qn("w:shd"))
            if shd is not None:
                body_fills.append(shd.get(qn("w:fill")))
ids = [r.cells[0].text.strip() for r in t.rows[1:]]
statuses = [r.cells[6].text.strip() for r in t.rows[1:]]
regression_ids = [r.cells[0].text.strip() for r in d.tables[1].rows[1:]]
print({"tables": len(d.tables), "main_columns": len(t.columns), "case_rows": len(t.rows)-1,
       "first_id": ids[0], "last_id": ids[-1], "statuses": {s: statuses.count(s) for s in sorted(set(statuses))},
       "regression_rows": len(regression_ids), "regression_ids": regression_ids,
       "body_cell_fills": body_fills, "execution_rows": len(d.tables[2].rows)-1})
