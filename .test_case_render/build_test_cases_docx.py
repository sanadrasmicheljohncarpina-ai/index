from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "TEST_CASES.md"
OUTPUT = ROOT / "docs" / "TEST_CASES_Run_2026-10-06.docx"
lines = SOURCE.read_text(encoding="utf-8").splitlines()

def parse_table(start):
    rows = []
    i = start
    while i < len(lines) and lines[i].startswith("|"):
        if i > start + 1:
            rows.append([cell.strip() for cell in lines[i].strip().strip("|").split("|")])
        i += 1
    return rows, i

cases = next(parse_table(i)[0] for i, line in enumerate(lines) if line.startswith("| Test Case ID |"))
regression_heading_index = next(i for i, line in enumerate(lines) if line == "## Reported Dean regression checks")
regression_table_index = next(i for i in range(regression_heading_index + 1, len(lines)) if lines[i].startswith("| Test Case ID |"))
regressions = parse_table(regression_table_index)[0]
record_index = next(i for i, line in enumerate(lines) if line.startswith("| Field | Value |"))
record = parse_table(record_index)[0]

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(.42)
section.bottom_margin = Inches(.42)
section.left_margin = Inches(.36)
section.right_margin = Inches(.36)

normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(7.2)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.0

heading = doc.add_paragraph()
heading.paragraph_format.space_after = Pt(5)
r = heading.add_run("Development and Testing")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(11)
r.font.color.rgb = RGBColor(0, 0, 0)

caption = doc.add_paragraph()
caption.paragraph_format.space_after = Pt(3)
r = caption.add_run("Test Case")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(8.5)
r.font.highlight_color = 7  # yellow highlighter, as in the supplied reference

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=48, start=48, bottom=48, end=48):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")

def set_borders(table, color="000000", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)

def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)

def fill_cell(cell, text, *, bold=False, size=7.2, align=WD_ALIGN_PARAGRAPH.LEFT, highlight=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    r.font.name = "Arial"
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0, 0, 0)
    r.bold = bold
    if highlight:
        r.font.highlight_color = 7

table = doc.add_table(rows=1, cols=7)
table.autofit = False
widths = [0.79, 1.05, 1.45, 1.07, 1.28, 1.36, 0.68]
headers = ["Test Case\nID", "Test Scenario", "Test Steps", "Test Data", "Expected Results", "Actual Results", "Passed/\nFailed"]
for j, (cell, label) in enumerate(zip(table.rows[0].cells, headers)):
    cell.width = Inches(widths[j])
    set_cell_shading(cell, "D9D9D9")
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell, top=60, bottom=60)
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    if j == 0:
        r = p.add_run("Test Case")
        r.font.highlight_color = 7
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(7.2)
        r2 = p.add_run("\nID")
        r2.bold = True
        r2.font.name = "Arial"
        r2.font.size = Pt(7.2)
    else:
        r = p.add_run(label)
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(7.2)
set_repeat_table_header(table.rows[0])

for row in cases:
    if len(row) != 7:
        raise ValueError(f"Expected 7 test columns, found {len(row)}: {row}")
    cells = table.add_row().cells
    for j, (cell, value) in enumerate(zip(cells, row)):
        cell.width = Inches(widths[j])
        align = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 6) else WD_ALIGN_PARAGRAPH.LEFT
        fill_cell(cell, value, bold=(j == 0), size=6.9, align=align)

set_borders(table)

regression_heading = doc.add_paragraph()
regression_heading.paragraph_format.space_before = Pt(7)
regression_heading.paragraph_format.space_after = Pt(3)
r = regression_heading.add_run("Reported Dean Regression Checks")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0, 0, 0)

regression_table = doc.add_table(rows=1, cols=7)
regression_table.autofit = False
regression_headers = ["Test Case\nID", "Test Scenario", "Test Steps", "Test Data", "Expected Results", "Actual Results", "Passed/\nFailed"]
for j, (cell, label) in enumerate(zip(regression_table.rows[0].cells, regression_headers)):
    cell.width = Inches(widths[j])
    set_cell_shading(cell, "D9D9D9")
    fill_cell(cell, label, bold=True, size=7.2,
              align=WD_ALIGN_PARAGRAPH.CENTER, highlight=(j == 0))
set_repeat_table_header(regression_table.rows[0])
for row in regressions:
    if len(row) != 7:
        raise ValueError(f"Expected 7 regression columns, found {len(row)}: {row}")
    cells = regression_table.add_row().cells
    for j, (cell, value) in enumerate(zip(cells, row)):
        cell.width = Inches(widths[j])
        align = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 6) else WD_ALIGN_PARAGRAPH.LEFT
        fill_cell(cell, value, bold=(j == 0), size=6.9, align=align)
set_borders(regression_table)

spacer = doc.add_paragraph()
spacer.paragraph_format.space_after = Pt(1)
record_heading = doc.add_paragraph()
record_heading.paragraph_format.space_before = Pt(5)
record_heading.paragraph_format.space_after = Pt(3)
r = record_heading.add_run("Execution Record")
r.bold = True
r.font.name = "Arial"
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0, 0, 0)

record_table = doc.add_table(rows=1, cols=2)
record_table.autofit = False
record_widths = [1.45, 6.23]
for i, head in enumerate(("Field", "Value")):
    cell = record_table.rows[0].cells[i]
    cell.width = Inches(record_widths[i])
    set_cell_shading(cell, "D9D9D9")
    fill_cell(cell, head, bold=True, size=7.2, align=WD_ALIGN_PARAGRAPH.LEFT)
set_repeat_table_header(record_table.rows[0])
for key, value in record:
    cells = record_table.add_row().cells
    for i, (cell, text) in enumerate(zip(cells, (key, value))):
        cell.width = Inches(record_widths[i])
        fill_cell(cell, text, bold=(i == 0), size=7.0)
set_borders(record_table)

doc.core_properties.title = "EPEMS Development and Testing Cases"
doc.core_properties.subject = "Executed EPEMS functional test cases and results"
doc.save(OUTPUT)
print(f"created={OUTPUT}")
print(f"case_rows={len(cases)} regression_rows={len(regressions)} execution_fields={len(record)}")
